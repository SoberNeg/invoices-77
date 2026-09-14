# -*- coding: utf-8 -*-
"""
Прототип: OCR счетов -> таблица закупок

Столбцы:
Покупатель | Поставщик | Дата | № Счета | Наименование товара | Кол-во

Поддержка:
- PDF (текстовый слой; при его отсутствии можно подключить PaddleOCR)
- XLSX
- XLS

Этот файл является отдельным прототипом и не связан с основной сборкой LED-калькулятора.
"""

from __future__ import annotations

import csv
import re
import threading
from dataclasses import dataclass
from pathlib import Path
from datetime import datetime
import tkinter as tk
from tkinter import ttk, filedialog, messagebox

# Drag & Drop — через tkinterdnd2.
# Если библиотека не установлена, программа продолжит работать обычной кнопкой.
try:
    from tkinterdnd2 import DND_FILES, TkinterDnD
    DND_AVAILABLE = True
except ImportError:
    DND_FILES = None
    TkinterDnD = None
    DND_AVAILABLE = False


APP_TITLE = "Прототип OCR счетов v5.1"

MONTHS = {
    "января": 1, "февраля": 2, "марта": 3, "апреля": 4,
    "мая": 5, "июня": 6, "июля": 7, "августа": 8,
    "сентября": 9, "октября": 10, "ноября": 11, "декабря": 12,
}


@dataclass
class RowData:
    position: str = ""
    buyer: str = ""
    supplier: str = ""
    date: str = ""
    invoice: str = ""
    item: str = ""
    qty: str = ""


def clean(s: str) -> str:
    return re.sub(r"\s+", " ", (s or "").replace("\xa0", " ")).strip()


def normalize_date(s: str) -> str:
    s = clean(s).lower().replace("г.", "").strip()

    m = re.search(r"(\d{1,2})[./-](\d{1,2})[./-](\d{4})", s)
    if m:
        d, mo, y = map(int, m.groups())
        try:
            return datetime(y, mo, d).strftime("%d.%m.%Y")
        except ValueError:
            return s

    m = re.search(
        r"(\d{1,2})\s+(" + "|".join(MONTHS) + r")\s+(\d{4})",
        s,
        flags=re.I,
    )
    if m:
        d = int(m.group(1))
        mo = MONTHS[m.group(2)]
        y = int(m.group(3))
        return datetime(y, mo, d).strftime("%d.%m.%Y")

    return s


def company_name(s: str) -> str:
    """Нормализует название организации/ИП без ИНН, КПП и адреса."""
    s = clean(s).replace("«", '"').replace("»", '"')

    s = re.sub(
        r"\bОБЩЕСТВО\s+С\s+ОГРАНИЧЕННОЙ\s+ОТВЕТСТВЕННОСТЬЮ\b",
        "ООО", s, flags=re.I
    )
    s = re.sub(
        r"\bОбщество\s+с\s+ограниченной\s+ответственностью\b",
        "ООО", s, flags=re.I
    )
    s = re.sub(
        r"\bИндивидуальный\s+предприниматель\b",
        "ИП", s, flags=re.I
    )

    # Реквизиты до названия.
    s = re.sub(r"^ИНН[: ]*\d+[,\s]*", "", s, flags=re.I)
    s = re.sub(r"^КПП[: ]*\d+[,\s]*", "", s, flags=re.I)
    s = re.sub(
        r"^ИНН\s+\d+[,\s]+КПП\s+\d+[,\s]+",
        "", s, flags=re.I
    )

    patterns = [
        r'(ИП\s+[А-ЯЁA-Z][^,;\n]+)',
        r'(ООО\s+(?:ТД\s+)?\"[^\"]+\")',
        r'(ООО\s+[А-ЯЁA-Z0-9][^,;\n]+)',
        r'(АО\s+\"[^\"]+\")',
        r'(АО\s+[А-ЯЁA-Z0-9][^,;\n]+)',
        r'(Интернет Решения,\s*ООО)',
        r'(П10\s+РУ\s+ООО)',
    ]
    for p in patterns:
        m = re.search(p, s, flags=re.I)
        if m:
            return clean(m.group(1)).strip(" ,;:")

    s = re.split(
        r",?\s+(?:ИНН|КПП|\d{6},|г\.|ул\.|пр-кт|дом\s|тел\.)",
        s, maxsplit=1, flags=re.I
    )[0]
    return clean(s).strip(" ,;:")

def extract_meta(text: str) -> dict:
    """Извлекает покупателя, поставщика, дату и номер счета."""
    flat = clean(text)
    lines = [clean(x) for x in text.splitlines() if clean(x)]

    invoice = ""
    date = ""

    invoice_patterns = [
        r'(?:Счет|Счёт|СЧЕТ)(?:-Оферта)?(?:\s+на\s+оплату)?\s*№\s*'
        r'([A-Za-zА-Яа-яЁё0-9._/-]+)\s+от\s+'
        r'([0-9]{1,2}(?:[./-][0-9]{1,2}[./-][0-9]{4}|\s+[А-Яа-яЁё]+\s+[0-9]{4})(?:\s*г\.?)?)',
        r'Оплата\s+по\s+заказу\s+([A-Za-zА-Яа-яЁё0-9._/-]+)\s+от\s+'
        r'([0-9]{1,2}[./-][0-9]{1,2}[./-][0-9]{4})',
    ]
    for p in invoice_patterns:
        m = re.search(p, flat, flags=re.I)
        if m:
            invoice = clean(m.group(1))
            date = normalize_date(m.group(2))
            break

    def after_label(labels, reject=""):
        for i, line in enumerate(lines):
            low = line.lower()
            if any(low.startswith(lbl) for lbl in labels):
                # Значение может быть на той же строке.
                raw = re.sub(
                    r"^(?:Поставщик|Покупатель|Плательщик|Получатель)"
                    r"(?:\s*\([^)]*\))?\s*:?\s*",
                    "", line, flags=re.I
                )
                candidates = []
                if raw and raw.lower() not in {"поставщик", "покупатель", "плательщик", "получатель"}:
                    candidates.append(raw)
                candidates.extend(lines[i + 1:i + 7])

                # Склеиваем соседние строки, если юр. название разорвано переносом.
                # Для строки с незакрытой кавычкой сначала пробуем склеенный вариант.
                joined = []
                for j, c in enumerate(candidates):
                    if j + 1 < len(candidates) and c.count('"') % 2 == 1:
                        joined.append(c + " " + candidates[j + 1])
                    joined.append(c)
                    if j + 1 < len(candidates) and c.count('"') % 2 == 0:
                        joined.append(c + " " + candidates[j + 1])

                for c in joined:
                    if reject and reject.lower() in c.lower():
                        continue
                    if (
                        re.search(r"\b(?:ООО|АО|ИП)\b", c, flags=re.I)
                        or re.search(r"Общество\s+с\s+ограниченной\s+ответственностью", c, flags=re.I)
                        or re.search(r"Индивидуальный\s+предприниматель", c, flags=re.I)
                        or "Интернет Решения" in c
                        or re.search(r"\bП10\s+РУ\s+ООО\b", c, flags=re.I)
                    ):
                        name = company_name(c)
                        if name and "банк" not in name.lower():
                            return name
        return ""

    supplier = after_label(["поставщик"])

    # Если поставщика нет, ищем получателя в банковском блоке.
    if not supplier:
        for i, line in enumerate(lines):
            if line.lower() == "получатель":
                for c in lines[max(0, i - 4):i + 5]:
                    if (
                        re.search(r"\b(?:ООО|АО|ИП)\b", c, flags=re.I)
                        or re.search(r"Общество\s+с\s+ограниченной\s+ответственностью", c, flags=re.I)
                        or re.search(r"Индивидуальный\s+предприниматель", c, flags=re.I)
                    ):
                        name = company_name(c)
                        if name and "банк" not in name.lower():
                            supplier = name
                            break
                if supplier:
                    break

    buyer = after_label(["покупатель", "плательщик"], reject=supplier)

    # Универсальные fallback'и.
    if not supplier and "Интернет Решения" in flat:
        supplier = "Интернет Решения, ООО"

    if not buyer:
        orgs = []
        for line in lines:
            if (
                re.search(r"\bИНН\b", line, flags=re.I)
                and (
                    re.search(r"\b(?:ООО|АО|ИП)\b", line, flags=re.I)
                    or re.search(r"Общество\s+с\s+ограниченной\s+ответственностью", line, flags=re.I)
                    or re.search(r"Индивидуальный\s+предприниматель", line, flags=re.I)
                )
            ):
                name = company_name(line)
                if name and "банк" not in name.lower():
                    orgs.append(name)
        for name in orgs:
            if clean(name).lower() != clean(supplier).lower():
                buyer = name
                break

    return {
        "buyer": buyer,
        "supplier": supplier,
        "date": date,
        "invoice": invoice,
    }

def extract_pdf_text(path: Path) -> str:
    try:
        import fitz
    except ImportError:
        raise RuntimeError("Установите PyMuPDF: pip install pymupdf")

    doc = fitz.open(path)
    return "\n".join(page.get_text("text") for page in doc)


def extract_pdf_text_with_ocr(path: Path, force: bool = False) -> str:
    text = extract_pdf_text(path)

    # Если в PDF есть нормальный текстовый слой — OCR не нужен, кроме force=True.
    if not force and len(re.findall(r"[A-Za-zА-Яа-яЁё]", text)) > 100:
        return text

    try:
        import fitz
        import numpy as np
        from PIL import Image
        from paddleocr import PaddleOCR
    except ImportError:
        raise RuntimeError(
            "PDF является сканом. Для OCR установите:\n"
            "pip install paddleocr paddlepaddle pillow numpy"
        )

    try:
        ocr = PaddleOCR(lang="ru")
    except Exception:
        ocr = PaddleOCR(lang="en")

    doc = fitz.open(path)
    result_lines = []

    for page in doc:
        pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
        img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
        arr = np.array(img)

        try:
            result = ocr.ocr(arr)
            for group in result or []:
                for line in group or []:
                    if len(line) >= 2:
                        result_lines.append(str(line[1][0]))
        except Exception:
            pass

    return "\n".join(result_lines)


def extract_items_generic(text: str) -> list[tuple[str, str]]:
    """Обычный случай: вся товарная строка находится в одной строке текста."""
    rows = []
    lines = [clean(x) for x in text.splitlines() if clean(x)]

    for line in lines:
        m = re.match(
            r"^\d+\s+(.+?)\s+(\d+(?:[.,]\d+)?)\s*(?:шт\.?|pcs\.?)\b",
            line,
            flags=re.I,
        )
        if m:
            rows.append((clean(m.group(1)), m.group(2).replace(",", ".")))

    return rows


def extract_items_vertical_table(text: str) -> list[tuple[str, str]]:
    """
    Макет №1: PDF-таблица читается PyMuPDF по колонкам.
    Пример: счет крепежа №9303.
    """
    pattern = re.compile(
        r"(?:^|\n)\s*(\d+(?:[.,]\d+)?)\s*\n"
        r"\s*(?:шт\.?|pcs\.?|уп\.?|упак\.?)\s*\n"
        r"\s*(.+?)\s*\n"
        r"\s*(\d{1,4})\s*\n"
        r"\s*([A-Za-zА-Яа-яЁё0-9._/-]{3,})\s*(?=\n|$)",
        flags=re.I | re.S,
    )

    rows = []
    for m in pattern.finditer(text):
        qty = m.group(1).replace(",", ".")
        name = clean(m.group(2))
        if name and not name.lower().startswith(("итого", "в том числе", "расчетный вес")):
            rows.append((name, qty))
    return rows


def extract_items_ledcapital(text: str) -> list[tuple[str, str]]:
    """
    Макет №2: LED Capital / ИП Мухамеджанов.
    Наименование начинается с номера позиции, количество находится на следующей строке.
    """
    lines = [clean(x) for x in text.splitlines() if clean(x)]
    rows = []

    for i, line in enumerate(lines):
        m = re.match(r"^(\d+)\s+(.+)$", line)
        if not m:
            continue

        # Отсекаем пункты текста вне товарной таблицы.
        if not re.search(r"[A-Za-zА-Яа-яЁё]", m.group(2)):
            continue

        if i + 1 < len(lines) and re.fullmatch(r"\d+(?:[.,]\d+)?", lines[i + 1]):
            # После количества должна быть единица "шт".
            if i + 2 < len(lines) and re.fullmatch(r"шт\.?", lines[i + 2], flags=re.I):
                rows.append((clean(m.group(2)), lines[i + 1].replace(",", ".")))

    return rows


def extract_items_tayle(text: str) -> list[tuple[str, str]]:
    """
    Макет №3: Тайле Рус.
    Позиция занимает много строк:
    № -> артикул (1-2 строки) -> описание (несколько строк) -> '4 шт'.
    """
    lines = [clean(x) for x in text.splitlines() if clean(x)]
    rows = []

    try:
        start = next(i for i, x in enumerate(lines) if x.lower() == "сумма") + 1
    except StopIteration:
        return rows

    # Ограничиваем товарный блок до "Итого".
    end = next(
        (i for i in range(start, len(lines)) if lines[i].lower().startswith("итого")),
        len(lines),
    )
    block = lines[start:end]

    # Индексы отдельных строк "1", "2", ...
    position_starts = [
        i for i, x in enumerate(block)
        if re.fullmatch(r"\d{1,3}", x)
    ]

    for k, pos in enumerate(position_starts):
        next_pos = position_starts[k + 1] if k + 1 < len(position_starts) else len(block)
        part = block[pos + 1:next_pos]

        qty_idx = -1
        qty = ""
        for j, x in enumerate(part):
            m = re.fullmatch(r"(\d+(?:[.,]\d+)?)\s*шт\.?", x, flags=re.I)
            if m:
                qty_idx = j
                qty = m.group(1).replace(",", ".")
                break

        if qty_idx < 0:
            continue

        before_qty = part[:qty_idx]
        if not before_qty:
            continue

        # Артикул идет первым и часто разбит переносом.
        # Описание ищем с первой строки, явно похожей на русское товарное описание.
        desc_start = None
        for j, x in enumerate(before_qty):
            if re.search(r"[А-Яа-яЁё]{3,}", x):
                desc_start = j
                break

        if desc_start is None:
            continue

        name = clean(" ".join(before_qty[desc_start:]))
        if name:
            rows.append((name, qty))

    return rows


def extract_items_ozon(text: str) -> list[tuple[str, str]]:
    """
    Макет №4: Ozon.
    Наименование многострочное, далее категория, цена, количество, единица.
    """
    lines = [clean(x) for x in text.splitlines() if clean(x)]
    rows = []

    try:
        start = next(i for i, x in enumerate(lines) if x.lower() == "сумма") + 1
    except StopIteration:
        return rows

    end = next(
        (i for i in range(start, len(lines)) if lines[i].lower().startswith("итого")),
        len(lines),
    )
    block = lines[start:end]

    # Сейчас счет Ozon может содержать несколько позиций.
    # Начало позиции: "1 <название>" либо отдельный "1".
    i = 0
    while i < len(block):
        m_inline = re.match(r"^(\d+)\s+(.+)$", block[i])
        m_only = re.fullmatch(r"\d+", block[i])

        if not (m_inline or m_only):
            i += 1
            continue

        if m_inline:
            name_parts = [m_inline.group(2)]
            i += 1
        else:
            name_parts = []
            i += 1

        # Идем до количества+шт. Цена может быть отдельной строкой перед количеством.
        qty = ""
        while i < len(block):
            if re.fullmatch(r"\d+(?:[.,]\d+)?", block[i]):
                # Если следующая строка "шт.", это количество.
                if i + 1 < len(block) and re.fullmatch(r"шт\.?", block[i + 1], flags=re.I):
                    qty = block[i].replace(",", ".")
                    i += 2
                    break

                # Число с двумя десятичными часто цена — не включаем в имя.
                if re.fullmatch(r"\d[\d ]*[,.]\d{2}", block[i]):
                    i += 1
                    continue

            # Категорию Ozon стараемся не включать в название.
            low = block[i].lower()
            if low in {"электроника и", "электротовары", "электроника и электротовары"}:
                i += 1
                continue

            name_parts.append(block[i])
            i += 1

        name = clean(" ".join(name_parts))
        if name and qty:
            rows.append((name, qty))

    return rows


TABLE_UNITS_RE = (
    r"(?:шт\.?|pcs\.?|м2|м²|м|пог\.?\s*м|куб\.?\s*м|рул\.?|"
    r"упак\.?|уп\.?|компл\.?|кг|л|пач\.?)"
)

def _is_probable_code_line(s: str) -> bool:
    s = clean(s)
    if not s or len(s) > 45:
        return False
    if re.search(r"[А-Яа-яЁё]{4,}", s):
        return False
    return bool(re.fullmatch(r"[A-Za-zА-Яа-яЁё0-9._/-]+", s))


def extract_items_sequential_table(text: str) -> list[tuple[str, str]]:
    """
    Универсальный парсер типовых российских счетов.
    Работает с PDF, где PyMuPDF читает таблицу последовательностью:
    № -> наименование -> количество/единица -> цена -> сумма.
    Поддерживает многострочные наименования, артикулы и колонки НДС.
    """
    lines = [clean(x) for x in text.splitlines() if clean(x)]
    if not lines:
        return []

    header_keys = (
        "товары (работы, услуги)",
        "наименование товара",
        "название товара или услуги",
        "товар (услуга)",
    )
    header_positions = [
        i for i, x in enumerate(lines)
        if any(k in x.lower() for k in header_keys)
    ]
    start = header_positions[0] + 1 if header_positions else 0

    end = next(
        (
            i for i in range(start, len(lines))
            if re.match(
                r"^(?:Итого|Всего\s+к\s+оплате|Всего\s+наименований)",
                lines[i], flags=re.I
            )
        ),
        len(lines),
    )
    block = lines[start:end]

    qty_num = r"\d[\d ]*(?:[.,]\d+)?"
    rows = []
    expected = 1
    i = 0

    while i < len(block):
        line = block[i]
        item_first = ""

        if line == str(expected):
            pass
        else:
            m_inline = re.match(r"^(\d{1,3})\s+(.+)$", line)
            if not (m_inline and int(m_inline.group(1)) == expected):
                i += 1
                continue
            item_first = clean(m_inline.group(2))

        name_parts = [item_first] if item_first else []
        qty = ""
        qty_idx = -1

        j = i + 1
        while j < len(block) and j <= i + 24:
            x = block[j]

            m_same = re.fullmatch(
                rf"({qty_num})\s+({TABLE_UNITS_RE})",
                x, flags=re.I
            )
            if m_same:
                qty = clean(m_same.group(1)).replace(" ", "").replace(",", ".")
                qty_idx = j
                break

            if re.fullmatch(qty_num, x):
                if j + 1 < len(block) and re.fullmatch(TABLE_UNITS_RE, block[j + 1], flags=re.I):
                    qty = x.replace(" ", "").replace(",", ".")
                    qty_idx = j
                    break

            # Только после проверки количества проверяем начало следующей позиции.
            next_inline = re.match(r"^(\d{1,3})\s+(.+)$", x)
            if qty_idx < 0 and (
                x == str(expected + 1)
                or (next_inline and int(next_inline.group(1)) == expected + 1)
            ):
                break

            name_parts.append(x)
            j += 1

        if not qty:
            i += 1
            continue

        cleaned = []
        for x in name_parts:
            if not x:
                continue
            low = x.lower()
            if low in {
                "№", "артикул", "товар", "товары (работы, услуги)",
                "кол-во", "количество", "ед.", "ед", "ед.изм", "ед. изм.",
                "цена", "сумма", "код", "ставка", "ндс", "сумма ндс",
                "без ндс"
            }:
                continue
            if re.fullmatch(r"\d{1,3}", x):
                continue
            if re.fullmatch(r"\d[\d ]*[,.]\d{2}\s*₽?", x):
                continue
            if re.fullmatch(r"(?:Без\s+НДС|\d{1,2}%?)", x, flags=re.I):
                continue
            # Отдельный артикул/код не добавляем к наименованию.
            if _is_probable_code_line(x) and len(name_parts) > 1:
                continue
            cleaned.append(x)

        name = clean(" ".join(cleaned))
        if name:
            rows.append((name, qty))
            expected += 1
            i = qty_idx + 1
        else:
            i += 1

    return rows


def merge_text_and_ocr(path: Path, text: str, meta: dict) -> tuple[str, dict]:
    """
    Если PDF является сканом ИЛИ текстовый слой потерял важные реквизиты,
    добавляем OCR-текст и повторно извлекаем метаданные.
    """
    letters = len(re.findall(r"[A-Za-zА-Яа-яЁё]", text))
    need_ocr = letters < 100 or not meta.get("buyer") or not meta.get("supplier")

    if not need_ocr:
        return text, meta

    try:
        ocr_text = extract_pdf_text_with_ocr(path)
    except Exception:
        return text, meta

    # extract_pdf_text_with_ocr для нормального текстового PDF может вернуть тот же текст.
    # Поэтому отдельно OCR принудительно делаем только если текста почти нет.
    if letters >= 100 and ocr_text == text:
        return text, meta

    combined = text + "\n" + ocr_text
    return combined, extract_meta(combined)

def parse_pdf(path: Path) -> list[RowData]:
    text = extract_pdf_text(path)
    meta = extract_meta(text)
    letters = len(re.findall(r"[A-Za-zА-Яа-яЁё]", text))

    # Для сканов OCR обязателен.
    if letters < 100:
        text = extract_pdf_text_with_ocr(path, force=True)
        meta = extract_meta(text)
    # Для некоторых PDF текстовый слой содержит таблицу, но не содержит
    # покупателя/поставщика (например, часть текста сохранена как изображение).
    elif not meta.get("buyer") or not meta.get("supplier"):
        try:
            ocr_text = extract_pdf_text_with_ocr(path, force=True)
            combined = text + "\n" + ocr_text
            ocr_meta = extract_meta(combined)
            if ocr_meta.get("buyer"):
                meta["buyer"] = ocr_meta["buyer"]
            if ocr_meta.get("supplier"):
                meta["supplier"] = ocr_meta["supplier"]
        except Exception:
            pass

    low = text.lower()
    items = []

    # Специализированные макеты.
    if "мухамеджанов" in low or "ledcapital" in low:
        items = extract_items_ledcapital(text)
    elif "тайле рус" in low:
        items = extract_items_tayle(text)
    elif "ozon" in low or "интернет решения" in low or "счет-оферта" in low:
        items = extract_items_ozon(text)
    elif "инструменты на горской" in low:
        items = extract_items_vertical_table(text)

    # Новый универсальный парсер типовых счетов.
    if not items:
        items = extract_items_sequential_table(text)

    # Старые fallback'и.
    if not items:
        items = extract_items_generic(text)
    if not items:
        items = extract_items_vertical_table(text)

    return [
        RowData(
            position=str(n),
            buyer=meta["buyer"],
            supplier=meta["supplier"],
            date=meta["date"],
            invoice=meta["invoice"],
            item=name,
            qty=qty,
        )
        for n, (name, qty) in enumerate(items, start=1)
    ]

def parse_xlsx(path: Path) -> list[RowData]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise RuntimeError("Установите openpyxl: pip install openpyxl")

    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]

    matrix = [
        ["" if v is None else str(v) for v in row]
        for row in ws.iter_rows(values_only=True)
    ]

    return parse_matrix(matrix)


def parse_xls(path: Path) -> list[RowData]:
    try:
        import xlrd
    except ImportError:
        raise RuntimeError("Установите xlrd: pip install xlrd")

    wb = xlrd.open_workbook(str(path))
    ws = wb.sheet_by_index(0)

    matrix = []
    for r in range(ws.nrows):
        row = []
        for c in range(ws.ncols):
            v = ws.cell_value(r, c)
            if isinstance(v, float) and v.is_integer():
                v = int(v)
            row.append("" if v is None else str(v))
        matrix.append(row)

    return parse_matrix(matrix)


def parse_matrix(matrix: list[list[str]]) -> list[RowData]:
    """
    Макет №5: XLS/XLSX.
    Не используем OCR — читаем реальные ячейки.

    Поддерживаем:
    - 1С-подобные счета с объединенными ячейками;
    - заголовки "Товар", "Товар (Услуга)", "Наименование";
    - "Кол-во" / "Количество".
    """
    normalized = [
        [clean("" if v is None else str(v)) for v in row]
        for row in matrix
    ]

    text = "\n".join(
        " | ".join(v for v in row if v)
        for row in normalized
    )

    meta = extract_meta(text)

    # В Excel надежнее отдельно искать реквизиты по строкам.
    for ri, row in enumerate(normalized):
        joined = " | ".join(v for v in row if v)
        low = joined.lower()

        # Счет.
        if "счет на оплату" in low or "счёт на оплату" in low:
            m = re.search(
                r"(?:Счет|Счёт)\s+на\s+оплату\s*№\s*([^\s|]+)\s+от\s+(.+?)(?:\||$)",
                joined,
                flags=re.I,
            )
            if m:
                meta["invoice"] = clean(m.group(1))
                meta["date"] = normalize_date(m.group(2))

        # Поставщик/покупатель могут находиться в другой ячейке той же строки.
        if any(v.lower().rstrip(":") == "поставщик" for v in row):
            candidates = [
                v for v in row
                if v and v.lower().rstrip(":") not in {"поставщик", "(исполнитель)"}
            ]
            if candidates:
                meta["supplier"] = company_name(max(candidates, key=len))

        if any(v.lower().rstrip(":") == "покупатель" for v in row):
            candidates = [
                v for v in row
                if v and v.lower().rstrip(":") not in {"покупатель", "(заказчик)"}
            ]
            if candidates:
                meta["buyer"] = company_name(max(candidates, key=len))

    # Ищем строку заголовка таблицы.
    header_row = -1
    name_col = -1
    qty_col = -1

    for ri, row in enumerate(normalized):
        local_name = -1
        local_qty = -1

        for ci, value in enumerate(row):
            low = value.lower()

            if (
                "наименование" in low
                or low == "товар"
                or low.startswith("товар ")
                or low.startswith("товар(")
                or "товар (услуга)" in low
                or "товары (работы, услуги)" in low
                or low == "продукт"
            ):
                local_name = ci

            if "кол-во" in low or "количество" in low or low == "кол.":
                local_qty = ci

        if local_name >= 0 and local_qty >= 0:
            header_row = ri
            name_col = local_name
            qty_col = local_qty
            break

    rows = []

    if header_row >= 0:
        for row in normalized[header_row + 1:]:
            if max(name_col, qty_col) >= len(row):
                continue

            name = clean(row[name_col])
            qty_cell = clean(row[qty_col])

            # Техническая строка 1|2|3|4... в печатной форме 1С.
            if name.isdigit() and qty_cell.isdigit():
                continue

            if not name:
                continue

            if name.lower().startswith(("итого", "всего", "ндс")):
                break

            q = re.search(r"\d+(?:[.,]\d+)?", qty_cell)
            if not q:
                continue

            rows.append(
                RowData(
                    position=str(len(rows) + 1),
                    buyer=meta["buyer"],
                    supplier=meta["supplier"],
                    date=meta["date"],
                    invoice=meta["invoice"],
                    item=name,
                    qty=q.group(0).replace(",", "."),
                )
            )

    return rows

def parse_file(path: Path) -> list[RowData]:
    ext = path.suffix.lower()

    if ext == ".pdf":
        return parse_pdf(path)

    if ext == ".xlsx":
        return parse_xlsx(path)

    if ext == ".xls":
        return parse_xls(path)

    raise RuntimeError(f"Формат {ext} не поддерживается")


BaseTk = TkinterDnD.Tk if DND_AVAILABLE else tk.Tk


class App(BaseTk):
    def __init__(self):
        super().__init__()

        self.title(APP_TITLE)
        self.geometry("1220x700")
        self.minsize(950, 560)

        self.files: list[Path] = []
        self.rows: list[RowData] = []
        self.loaded_count = tk.StringVar(value="Загружено счетов: 0")

        self.build_ui()

    def build_ui(self):
        header = ttk.Frame(self, padding=12)
        header.pack(fill="x")

        ttk.Label(
            header,
            text="OCR счетов → таблица",
            font=("Arial", 17, "bold"),
        ).pack(side="left")

        controls = ttk.Frame(self, padding=(12, 0))
        controls.pack(fill="x")

        ttk.Button(
            controls,
            text="Добавить счета",
            command=self.add_files,
        ).pack(side="left")

        ttk.Button(
            controls,
            text="Распознать",
            command=self.start_parse,
        ).pack(side="left", padx=6)

        ttk.Button(
            controls,
            text="Очистить",
            command=self.clear_all,
        ).pack(side="left")

        ttk.Button(
            controls,
            text="Экспорт XLSX",
            command=self.export_xlsx,
        ).pack(side="right")

        # ------------------------------------------------------------
        # Блок Drag & Drop
        # ------------------------------------------------------------
        drop_outer = ttk.Frame(self, padding=(12, 10))
        drop_outer.pack(fill="x")

        self.drop_frame = tk.Frame(
            drop_outer,
            height=92,
            bd=2,
            relief="groove",
            bg="#f2f2f2",
        )
        self.drop_frame.pack(fill="x")
        self.drop_frame.pack_propagate(False)

        self.drop_label = tk.Label(
            self.drop_frame,
            text=(
                "Перетащите сюда файлы счетов PDF / XLSX / XLS\n"
                "или нажмите сюда для выбора файлов"
            ),
            bg="#f2f2f2",
            fg="#333333",
            font=("Arial", 12),
            cursor="hand2",
            justify="center",
        )
        self.drop_label.pack(fill="both", expand=True)

        self.drop_label.bind("<Button-1>", lambda _e: self.add_files())
        self.drop_frame.bind("<Button-1>", lambda _e: self.add_files())

        if DND_AVAILABLE:
            self.drop_frame.drop_target_register(DND_FILES)
            self.drop_frame.dnd_bind("<<Drop>>", self.on_drop_files)

            self.drop_label.drop_target_register(DND_FILES)
            self.drop_label.dnd_bind("<<Drop>>", self.on_drop_files)
        else:
            self.drop_label.configure(
                text=(
                    "Перетащите сюда файлы счетов PDF / XLSX / XLS\n"
                    "Drag & Drop требует tkinterdnd2 — пока можно нажать сюда"
                )
            )

        file_title = ttk.Frame(self, padding=(12, 0))
        file_title.pack(fill="x")

        ttk.Label(
            file_title,
            text="Добавленные файлы:",
            font=("Arial", 11, "bold"),
        ).pack(side="left")

        counter_frame = ttk.Frame(file_title, padding=(10, 4))
        counter_frame.pack(side="right")

        ttk.Label(
            counter_frame,
            textvariable=self.loaded_count,
            font=("Arial", 11, "bold"),
        ).pack()

        self.files_box = tk.Listbox(self, height=5)
        self.files_box.pack(fill="x", padx=12, pady=(4, 10))

        self.progress = ttk.Progressbar(self, mode="determinate", maximum=100)
        self.progress.pack(fill="x", padx=12)

        self.status = tk.StringVar(value="Готово")
        ttk.Label(self, textvariable=self.status).pack(anchor="w", padx=12, pady=(4, 8))

        frame = ttk.Frame(self)
        frame.pack(fill="both", expand=True, padx=12, pady=(0, 12))

        cols = ("position", "buyer", "supplier", "date", "invoice", "item", "qty")

        self.tree = ttk.Treeview(frame, columns=cols, show="headings")
        self.tree.grid(row=0, column=0, sticky="nsew")

        names = {
            "position": "№ п/п",
            "buyer": "Покупатель",
            "supplier": "Поставщик",
            "date": "Дата",
            "invoice": "№ Счета",
            "item": "Наименование товара",
            "qty": "Кол-во",
        }

        widths = {
            "position": 60,
            "buyer": 170,
            "supplier": 220,
            "date": 100,
            "invoice": 140,
            "item": 470,
            "qty": 90,
        }

        for col in cols:
            self.tree.heading(col, text=names[col])
            self.tree.column(col, width=widths[col], anchor="w")

        y = ttk.Scrollbar(frame, orient="vertical", command=self.tree.yview)
        y.grid(row=0, column=1, sticky="ns")

        x = ttk.Scrollbar(frame, orient="horizontal", command=self.tree.xview)
        x.grid(row=1, column=0, sticky="ew")

        self.tree.configure(yscrollcommand=y.set, xscrollcommand=x.set)

        frame.rowconfigure(0, weight=1)
        frame.columnconfigure(0, weight=1)

    def add_paths(self, paths):
        """Добавляет пути из диалога или Drag & Drop."""
        allowed = {".pdf", ".xlsx", ".xls"}

        added = 0
        skipped = []

        for raw in paths:
            if not raw:
                continue

            p = Path(raw)

            if p.suffix.lower() not in allowed:
                skipped.append(p.name)
                continue

            if not p.exists() or not p.is_file():
                skipped.append(p.name)
                continue

            if p not in self.files:
                self.files.append(p)
                number = len(self.files)
                self.files_box.insert("end", f"{number}.  {p}")
                added += 1

        self.loaded_count.set(f"Загружено счетов: {len(self.files)}")

        if added:
            self.status.set(f"Добавлено файлов: {added}. Всего: {len(self.files)}")

        if skipped:
            messagebox.showwarning(
                "Некоторые файлы пропущены",
                "Поддерживаются только PDF, XLSX и XLS.\n\n"
                + "\n".join(skipped[:10])
            )

    def on_drop_files(self, event):
        """Обрабатывает файлы, перетащенные в блок."""
        try:
            # splitlist правильно разбирает пути с пробелами и фигурными скобками.
            dropped = self.tk.splitlist(event.data)
            self.add_paths(dropped)

            # Короткая визуальная реакция на успешный drop.
            self.drop_frame.configure(bg="#e8f5e9")
            self.drop_label.configure(bg="#e8f5e9")
            self.after(
                500,
                lambda: (
                    self.drop_frame.configure(bg="#f2f2f2"),
                    self.drop_label.configure(bg="#f2f2f2"),
                ),
            )
        except Exception as e:
            messagebox.showerror("Drag & Drop", f"Не удалось добавить файл:\n{e}")

    def add_files(self):
        files = filedialog.askopenfilenames(
            title="Выберите счета",
            filetypes=[
                ("Счета", "*.pdf *.xlsx *.xls"),
                ("PDF", "*.pdf"),
                ("Excel", "*.xlsx *.xls"),
            ],
        )
        self.add_paths(files)

    def clear_all(self):
        self.files.clear()
        self.rows.clear()
        self.files_box.delete(0, "end")
        self.loaded_count.set("Загружено счетов: 0")

        for iid in self.tree.get_children():
            self.tree.delete(iid)

        self.progress["value"] = 0
        self.status.set("Готово")

    def start_parse(self):
        if not self.files:
            messagebox.showwarning("Нет файлов", "Добавьте хотя бы один счет.")
            return

        self.rows.clear()

        for iid in self.tree.get_children():
            self.tree.delete(iid)

        threading.Thread(target=self.worker, daemon=True).start()

    def worker(self):
        try:
            total = len(self.files)

            for i, path in enumerate(self.files, start=1):
                self.after(0, lambda p=path: self.status.set(f"Обработка: {p.name}"))

                rows = parse_file(path)
                self.rows.extend(rows)

                self.after(
                    0,
                    lambda value=(i / total) * 100: self.progress.configure(value=value),
                )

                self.after(0, self.refresh)

            self.after(
                0,
                lambda: self.status.set(f"Готово. Строк: {len(self.rows)}"),
            )

        except Exception as e:
            self.after(0, lambda: messagebox.showerror("Ошибка", str(e)))
            self.after(0, lambda: self.status.set("Ошибка"))

    def refresh(self):
        for iid in self.tree.get_children():
            self.tree.delete(iid)

        for row in self.rows:
            self.tree.insert(
                "",
                "end",
                values=(
                    row.position,
                    row.buyer,
                    row.supplier,
                    row.date,
                    row.invoice,
                    row.item,
                    row.qty,
                ),
            )

    def export_xlsx(self):
        if not self.rows:
            messagebox.showwarning("Нет данных", "Сначала распознайте счета.")
            return

        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
            from openpyxl.utils import get_column_letter
        except ImportError:
            messagebox.showerror("Ошибка", "Установите openpyxl")
            return

        filename = filedialog.asksaveasfilename(
            title="Сохранить таблицу",
            defaultextension=".xlsx",
            initialfile="таблица_счетов.xlsx",
            filetypes=[("Excel", "*.xlsx")],
        )

        if not filename:
            return

        wb = Workbook()
        ws = wb.active
        ws.title = "Счета"

        headers = [
            "№ п/п",
            "Покупатель",
            "Поставщик",
            "Дата",
            "№ Счета",
            "Наименование товара",
            "Кол-во",
            "ИТОГО",
            "Остаток",
        ]

        ws.append(headers)

        # Вторая строка под шапкой — как в файле-образце.
        # В образце надпись расположена во второй колонке.
        ws.append(["", "Количество изделий", "", "", "", "", "", "", ""])

        for row in self.rows:
            ws.append([
                row.position,
                row.buyer,
                row.supplier,
                row.date,
                row.invoice,
                row.item,
                row.qty,
                "",
                "",
            ])

            # Одна пустая строка после каждой строки с данными.
            ws.append(["", "", "", "", "", "", "", "", ""])

        # ------------------------------------------------------------
        # Оформление шапки по образцу "Калькуляция Примыкание"
        # ------------------------------------------------------------
        # Светло-желтая заливка, черный жирный Times New Roman,
        # центрирование, перенос текста и черные границы.
        header_fill = PatternFill(
            fill_type="solid",
            fgColor="FFFFCC",
        )

        thin_black = Side(style="thin", color="000000")
        medium_black = Side(style="medium", color="000000")

        header_cells = ws[1]

        for index, cell in enumerate(header_cells, start=1):
            cell.font = Font(
                name="Times New Roman",
                size=12,
                bold=True,
                color="000000",
            )
            cell.fill = header_fill
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

            # Внешняя рамка шапки более толстая,
            # внутренние вертикальные линии — тонкие.
            left_side = medium_black if index == 1 else thin_black
            right_side = medium_black if index == len(header_cells) else thin_black

            cell.border = Border(
                left=left_side,
                right=right_side,
                top=medium_black,
                bottom=medium_black,
            )

        # Высота шапки близка к образцу, чтобы длинные названия
        # не сжимались в одну строку.
        ws.row_dimensions[1].height = 42

        # ------------------------------------------------------------
        # Строка под шапкой — "Количество изделий"
        # Вся рабочая ширина таблицы A:I заливается синим цветом,
        # как в файле-образце.
        # ------------------------------------------------------------
        quantity_fill = PatternFill(
            fill_type="solid",
            fgColor="00B0F0",
        )

        for col in range(1, 10):  # A:I
            cell = ws.cell(row=2, column=col)
            cell.fill = quantity_fill
            cell.font = Font(
                name="Times New Roman",
                size=12,
                bold=True,
                color="000000",
            )
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

            left_side = medium_black if col == 1 else thin_black
            right_side = medium_black if col == 9 else thin_black

            cell.border = Border(
                left=left_side,
                right=right_side,
                top=thin_black,
                bottom=medium_black,
            )

        ws["B2"] = "Количество изделий"
        ws.row_dimensions[2].height = 22

        widths = [8, 24, 30, 14, 18, 72, 12, 14, 14]

        for i, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = width

        # ------------------------------------------------------------
        # Границы всей рабочей таблицы
        # ------------------------------------------------------------
        # Все ячейки рабочего поля, включая пустые строки между позициями,
        # получают тонкие черные границы.
        #
        # Шапка (строка 1) остается с более толстыми границами,
        # заданными выше.
        data_border = Border(
            left=thin_black,
            right=thin_black,
            top=thin_black,
            bottom=thin_black,
        )

        for row_cells in ws.iter_rows(
            min_row=2,
            max_row=ws.max_row,
            min_col=1,
            max_col=9,
        ):
            for cell in row_cells:
                cell.border = data_border

        # Ячейка сразу под заголовком "ИТОГО" (H2):
        # только она имеет желтую заливку, как в образце.
        ws["H2"].fill = PatternFill(
            fill_type="solid",
            fgColor="FFFF00",
        )

        # Повторно задаем жирную рамку основной шапке A1:I1,
        # чтобы общий цикл оформления ее не затронул.
        for index, cell in enumerate(ws[1], start=1):
            left_side = medium_black if index == 1 else thin_black
            right_side = medium_black if index == 9 else thin_black

            cell.border = Border(
                left=left_side,
                right=right_side,
                top=medium_black,
                bottom=medium_black,
            )

        ws.freeze_panes = "A2"
        wb.save(filename)

        messagebox.showinfo("Готово", f"Файл сохранен:\n{filename}")


if __name__ == "__main__":
    App().mainloop()
